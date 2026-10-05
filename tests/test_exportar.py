import unittest
from datetime import date
from io import BytesIO
from types import SimpleNamespace as NS

from openpyxl import load_workbook

from tests import helpers  # noqa: F401
import exportar


def socio(**kw):
    base = dict(nombre="Ana", apellidos="Pérez", dni="12345678Z", direccion="Rúa Nova 3",
                fecha_nacimiento=date(1980, 5, 2), telefono="+34 600 111 222", email="a@b.es")
    return NS(**{**base, **kw})


class Exportar(unittest.TestCase):
    def hoja(self, socios, **kw):
        return load_workbook(BytesIO(exportar.construir_excel_socios(socios, **kw))).active

    def test_columnas_y_datos(self):
        ws = self.hoja([socio()])
        self.assertEqual(
            [c.value for c in ws[1]],
            ["Nombre", "Apellidos", "DNI / NIE", "Dirección", "Fecha de nacimiento", "Teléfono", "Correo electrónico"],
        )
        self.assertEqual([c.value for c in ws[2]][:4], ["Ana", "Pérez", "12345678Z", "Rúa Nova 3"])
        self.assertEqual(ws["F2"].value, "+34 600 111 222")          # texto, conserva el "+"
        self.assertEqual(ws["E2"].number_format, "DD/MM/YYYY")

    def test_un_dato_nunca_se_interpreta_como_formula(self):
        ws = self.hoja([socio(nombre="=CMD()")])
        self.assertEqual(ws["A2"].data_type, "s")

    def test_vacios_y_lista_vacia(self):
        ws = self.hoja([socio(direccion=None, email=None, fecha_nacimiento=None)])
        self.assertIsNone(ws["D2"].value)
        ws = self.hoja([])
        self.assertEqual((ws.max_row, ws.auto_filter.ref), (1, "A1:G1"))

    def test_color_de_cabecera(self):
        ws = self.hoja([socio()], color_cabecera="#17343f", color_texto="#ffffff")
        self.assertTrue(ws["A1"].fill.start_color.rgb.endswith("17343F"))

    def test_cabecera_fija_y_filtros(self):
        ws = self.hoja([socio(), socio(nombre="Luis")])
        self.assertEqual((ws.freeze_panes, ws.auto_filter.ref), ("A2", "A1:G3"))


if __name__ == "__main__":
    unittest.main()
