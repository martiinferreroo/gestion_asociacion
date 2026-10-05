"""Generación de la hoja de cálculo con los datos de los socios."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FUENTE = "Arial"

# (título de la columna, atributo del socio, ancho)
COLUMNAS = [
    ("Nombre", "nombre", 18),
    ("Apellidos", "apellidos", 26),
    ("DNI / NIE", "dni", 14),
    ("Dirección", "direccion", 40),
    ("Fecha de nacimiento", "fecha_nacimiento", 20),
    ("Teléfono", "telefono", 16),
    ("Correo electrónico", "email", 32),
]


def _hex(color: str) -> str:
    return color.lstrip("#").upper()


def construir_excel_socios(
    socios, color_cabecera="#0f766e", color_texto="#ffffff"
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Socios"

    relleno = PatternFill("solid", start_color=_hex(color_cabecera))
    for i, (titulo, _campo, ancho) in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=i, value=titulo)
        celda.font = Font(name=FUENTE, size=10, bold=True, color=_hex(color_texto))
        celda.fill = relleno
        celda.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[1].height = 22

    fila = 1
    for fila, socio in enumerate(socios, start=2):
        for i, (_titulo, campo, _ancho) in enumerate(COLUMNAS, start=1):
            valor = getattr(socio, campo)
            celda = ws.cell(row=fila, column=i)
            if valor is not None:
                celda.value = valor
                if isinstance(valor, str) and valor.startswith("="):
                    # Que un dato nunca se interprete como fórmula.
                    celda.data_type = "s"
            celda.font = Font(name=FUENTE, size=10)
            celda.alignment = Alignment(horizontal="left", vertical="center")
            if campo == "fecha_nacimiento":
                celda.number_format = "DD/MM/YYYY"

    ultima_col = get_column_letter(len(COLUMNAS))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ultima_col}{max(fila, 1)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
