"""Importación de socios desde Excel o CSV.

Este módulo no depende de la web ni de la base de datos: lee el archivo,
adivina qué hay en cada columna, limpia los datos y decide qué crear o
modificar (planificar). Quien lo usa se limita a ejecutar el plan.
"""
import csv
import io
import re
import unicodedata
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

DNI_GENERICO = "12345678Z"
MAX_FILAS = 5000
MAX_BYTES = 5 * 1024 * 1024
EXTENSIONES = (".xlsx", ".xlsm", ".csv")

CAMPOS_DESTINO = [
    ("", "No importar esta columna"),
    ("num_socio", "Nº de socio"),
    ("nombre", "Nombre"),
    ("apellidos", "Apellidos"),
    ("nombre_completo", "Nombre y apellidos juntos"),
    ("dni", "DNI / NIE"),
    ("fecha_nacimiento", "Fecha de nacimiento"),
    ("telefono", "Teléfono"),
    ("email", "Correo electrónico"),
    ("direccion", "Dirección"),
    ("fecha_alta", "Fecha de alta"),
    ("notas", "Observaciones"),
]
CLAVES = {k for k, _ in CAMPOS_DESTINO if k}

CAMPOS_DATOS = [
    "nombre", "apellidos", "dni", "fecha_nacimiento", "telefono", "email",
    "direccion", "fecha_alta", "notas",
]
ETIQUETA_CAMPO = {
    "nombre": "Nombre", "apellidos": "Apellidos", "dni": "DNI",
    "fecha_nacimiento": "Nacimiento", "telefono": "Teléfono",
    "email": "Correo", "direccion": "Dirección", "fecha_alta": "Alta",
    "notas": "Observaciones",
}

VACIOS = {
    "", "-", "--", "---", ".", "?", "n/a", "na", "nan", "none", "null", "s/d",
    "sd", "nd", "sin dato", "sin datos", "no tiene", "no consta", "ninguno",
}


def letra(indice: int) -> str:
    return get_column_letter(indice + 1)


# --------------------------------------------------------------------------
# Lectura del archivo
# --------------------------------------------------------------------------

def texto(v) -> str:
    """Valor de celda como texto limpio; '' si está vacío o es un marcador."""
    if v is None or isinstance(v, bool):
        return ""
    if isinstance(v, float):
        if v != v:
            return ""
        s = str(int(v)) if v.is_integer() else repr(v)
    elif isinstance(v, (datetime, date)):
        s = v.strftime("%d/%m/%Y")
    else:
        s = str(v)
    s = re.sub(r"\s+", " ", s.replace("\xa0", " ").replace("\u200b", "")).strip()
    return "" if s.lower() in VACIOS else s


def es_vacio(v) -> bool:
    return texto(v) == ""


def _recortar(filas):
    """Iguala el ancho de las filas y quita filas y columnas vacías del final."""
    while filas and all(es_vacio(c) for c in filas[-1]):
        filas.pop()
    ancho = 0
    for f in filas:
        for i, c in enumerate(f):
            if not es_vacio(c):
                ancho = max(ancho, i + 1)
    return [(list(f) + [None] * ancho)[:ancho] for f in filas]


def _leer_csv(ruta: Path):
    datos = ruta.read_bytes()
    contenido = None
    for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            contenido = datos.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    muestra = contenido[:4096]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=";,\t|")
    except csv.Error:
        dialecto = csv.excel
        dialecto.delimiter = ";" if muestra.count(";") > muestra.count(",") else ","
    return [list(f) for f in csv.reader(io.StringIO(contenido), dialecto)]


def cargar(ruta: Path, hoja=None):
    """Devuelve (nombres_de_hojas, hoja_usada, filas, recortado)."""
    ruta = Path(ruta)
    if ruta.suffix.lower() == ".csv":
        filas = _leer_csv(ruta)
        recortado = len(filas) > MAX_FILAS + 1
        return [], "", _recortar(filas[: MAX_FILAS + 1]), recortado

    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        hojas = list(wb.sheetnames)
        usada = hoja if hoja in hojas else hojas[0]
        filas, recortado = [], False
        for i, fila in enumerate(wb[usada].iter_rows(values_only=True)):
            if i > MAX_FILAS:
                recortado = True
                break
            filas.append(list(fila))
    finally:
        wb.close()
    return hojas, usada, _recortar(filas), recortado


# --------------------------------------------------------------------------
# Limpieza de datos
# --------------------------------------------------------------------------

PARTICULAS = {"de", "del", "la", "las", "los", "y", "e", "da", "do", "dos", "das", "van", "von", "di"}


def capitalizar(s: str) -> str:
    """Solo corrige nombres todo en mayúsculas o todo en minúsculas."""
    if not s or not (s.isupper() or s.islower()):
        return s
    s = re.sub(
        r"(^|[\s\-'’.])([^\W\d_])",
        lambda m: m.group(1) + m.group(2).upper(),
        s.lower(),
    )
    return " ".join(
        p.lower() if i > 0 and p.lower() in PARTICULAS else p
        for i, p in enumerate(s.split(" "))
    )


def dividir_nombre_completo(t: str):
    """'Pérez Ruiz, Ana' -> (Ana, Pérez Ruiz); 'Ana Pérez Ruiz' -> (Ana, Pérez Ruiz)."""
    t = capitalizar(t)
    if "," in t:
        apellidos, nombre = t.split(",", 1)
        return nombre.strip(), apellidos.strip()
    palabras = t.split()
    if len(palabras) <= 1:
        return t, ""
    i, apellidos, cuenta = len(palabras) - 1, [], 0
    while i > 0 and cuenta < 2:
        j = i
        while j - 1 > 0 and palabras[j - 1].lower() in PARTICULAS:
            j -= 1
        apellidos = palabras[j : i + 1] + apellidos
        i, cuenta = j - 1, cuenta + 1
    return " ".join(palabras[: i + 1]), " ".join(apellidos)


LETRAS_DNI = "TRWAGMYFPDXBNJZSQVHLCKE"


def limpiar_dni(v):
    """Devuelve (dni | None, mensaje | None)."""
    t = texto(v)
    if not t:
        return None, None
    s = re.sub(r"[\s.\-_/]", "", t).upper()
    if s == DNI_GENERICO:
        return None, None
    m = re.fullmatch(r"([XYZ])(\d{7})", s)
    if m:
        p, n = m.groups()
        dni = p + n + LETRAS_DNI[int(str("XYZ".index(p)) + n) % 23]
        return dni, f"Se añadió la letra que faltaba al NIE: {dni}."
    if re.fullmatch(r"\d{7,8}", s):
        n = s.zfill(8)
        dni = n + LETRAS_DNI[int(n) % 23]
        return dni, f"Se añadió la letra que faltaba al DNI: {dni}."
    if re.fullmatch(r"\d{7}[A-Z]", s):
        return "0" + s, None
    if re.fullmatch(r"[A-Z0-9]{5,15}", s) and re.search(r"\d", s):
        return s, None
    return None, f"DNI/NIE no válido: «{t}»."


def _formatear_telefono(solo: str) -> str:
    if re.fullmatch(r"\d{9}", solo):
        return f"{solo[:3]} {solo[3:6]} {solo[6:]}"
    m = re.fullmatch(r"\+34(\d{9})", solo)
    if m:
        d = m.group(1)
        return f"+34 {d[:3]} {d[3:6]} {d[6:]}"
    return solo


def separar_telefonos(v):
    """(válidos, inválidos). Admite varios números separados por / ; , y o."""
    t = texto(v)
    validos, invalidos = [], []
    if not t:
        return validos, invalidos
    for parte in re.split(r"\s*(?:[/;,]|\by\b|\bo\b)\s*", t):
        parte = re.sub(r"\.0+$", "", parte.strip())
        if not parte:
            continue
        solo = re.sub(r"[^\d+]", "", parte)
        if solo.startswith("00"):
            solo = "+" + solo[2:]
        if 6 <= len(solo.replace("+", "")) <= 15:
            f = _formatear_telefono(solo)
            if f not in validos:
                validos.append(f)
        else:
            invalidos.append(parte)
    return validos, invalidos


EMAIL_RE = re.compile(r"^[^@\s,;<>]+@[^@\s,;<>]+\.[^@\s,;<>.]+$")


def separar_emails(v):
    t = texto(v).lower().replace("mailto:", "")
    if not t:
        return [], []
    validos = []
    for parte in re.split(r"[;,\s]+", t):
        p = parte.strip("<>.'\"()")
        if p and EMAIL_RE.match(p) and p not in validos:
            validos.append(p)
    return (validos, []) if validos else ([], [t])


MESES = {
    "ene": 1, "jan": 1, "feb": 2, "mar": 3, "abr": 4, "apr": 4, "may": 5,
    "jun": 6, "jul": 7, "ago": 8, "aug": 8, "sep": 9, "set": 9, "oct": 10,
    "nov": 11, "dic": 12, "dec": 12,
}


def _sin_acentos(s: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c)
    )


def _anio(a: int, nacimiento: bool, hoy: date) -> int:
    if a >= 100:
        return a
    if nacimiento:
        return 2000 + a if a <= (hoy.year % 100) - 16 else 1900 + a
    return 2000 + a if 2000 + a <= hoy.year else 1900 + a


def _construir(d, m, a, hoy, nacimiento):
    if m > 12 and d <= 12:
        d, m = m, d
    try:
        f = date(_anio(a, nacimiento, hoy), m, d)
    except ValueError:
        return None
    return f if 1900 <= f.year and f <= hoy else None


def _fecha_texto(s, nacimiento, hoy):
    s = re.sub(r"[t\s]+\d{1,2}:\d{2}.*$", "", s.lower().strip())
    m = re.fullmatch(r"(\d{1,4})[/\-. ](\d{1,2})[/\-. ](\d{1,4})", s)
    if m:
        a, b, c = m.groups()
        if len(a) == 4:
            return _construir(int(c), int(b), int(a), hoy, nacimiento)
        return _construir(int(a), int(b), int(c), hoy, nacimiento)
    if re.fullmatch(r"\d{7,8}", s):
        s = s.zfill(8)
        if 1900 <= int(s[:4]) <= hoy.year and 1 <= int(s[4:6]) <= 12:
            return _construir(int(s[6:]), int(s[4:6]), int(s[:4]), hoy, nacimiento)
        return _construir(int(s[:2]), int(s[2:4]), int(s[4:]), hoy, nacimiento)
    if re.fullmatch(r"\d{6}", s):
        return _construir(int(s[:2]), int(s[2:4]), int(s[4:]), hoy, nacimiento)
    m = re.fullmatch(
        r"(\d{1,2})\s*(?:de\s+|[-/.]\s*)?([a-zñáéíóú]{3,})\.?\s*(?:de\s+|[-/.]\s*)?(\d{2,4})", s
    )
    if m:
        mes = MESES.get(_sin_acentos(m.group(2))[:3])
        if mes:
            return _construir(int(m.group(1)), mes, int(m.group(3)), hoy, nacimiento)
    return None


def parsear_fecha(v, nacimiento=True, hoy=None):
    """Acepta fechas de Excel, números de serie y textos en muchos formatos."""
    hoy = hoy or date.today()
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, datetime):
        f = v.date()
    elif isinstance(v, date):
        f = v
    elif isinstance(v, (int, float)):
        n = float(v)
        if 2101 <= n <= 80000:
            f = (datetime(1899, 12, 30) + timedelta(days=int(n))).date()
        elif n == int(n) and 1_000_000 <= n <= 99_999_999:
            return _fecha_texto(str(int(n)), nacimiento, hoy)
        else:
            return None
    else:
        return _fecha_texto(texto(v), nacimiento, hoy)
    return f if 1900 <= f.year and f <= hoy else None


# --------------------------------------------------------------------------
# Adivinar qué hay en cada columna
# --------------------------------------------------------------------------

def _normalizar(s: str) -> str:
    s = _sin_acentos(s.lower())
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]+", " ", s)).strip()


def adivinar_por_encabezado(encabezado: str) -> str:
    n = _normalizar(encabezado)
    if not n:
        return ""
    if re.search(r"nombre y apellidos|nombre completo|apellidos y nombre|apellidos nombre|nombre apellidos|titular", n):
        return "nombre_completo"
    if "apellid" in n or n in ("surname", "last name"):
        return "apellidos"
    if re.search(r"\b(dni|nif|nie|documento|cif)\b", n) or "identific" in n:
        return "dni"
    if "nacim" in n or n in ("fnac", "f nac", "fecha nac", "cumpleanos"):
        return "fecha_nacimiento"
    if re.search(r"alta|ingreso|inscrip", n):
        return "fecha_alta"
    if re.search(r"telef|movil|tlf|tfno|\btel\b|phone|celular", n):
        return "telefono"
    if re.search(r"mail|correo", n):
        return "email"
    if re.search(r"direcc|domic|calle|address", n):
        return "direccion"
    if re.search(r"observ|nota|coment", n):
        return "notas"
    if re.search(r"\bsocio\b|\bnumero\b|\bnum\b|^n$|^no$|^id$|codigo", n) and "nombre" not in n:
        return "num_socio"
    if "nombre" in n or n == "name":
        return "nombre"
    return ""


def _parece_fecha(v) -> bool:
    if isinstance(v, (datetime, date)):
        return True
    t = texto(v)
    return bool(t and re.search(r"\d[/\-. ]\d|[a-zA-Z]{3}", t) and parsear_fecha(t) is not None)


def adivinar_por_contenido(valores) -> str:
    muestra = [v for v in valores if not es_vacio(v)][:40]
    if not muestra:
        return ""

    def frac(pred):
        return sum(1 for v in muestra if pred(v)) / len(muestra)

    def compacto(v):
        return re.sub(r"[\s.\-]", "", texto(v)).upper()

    if frac(lambda v: "@" in texto(v)) >= 0.6:
        return "email"
    if frac(lambda v: re.fullmatch(r"[XYZ]?\d{7,8}[A-Z]", compacto(v))) >= 0.6:
        return "dni"
    if frac(lambda v: re.fullmatch(r"\+?\d{9,13}", re.sub(r"[\s.\-]", "", texto(v)))) >= 0.6:
        return "telefono"
    if frac(_parece_fecha) >= 0.6:
        return "fecha_nacimiento"
    if frac(lambda v: re.fullmatch(r"\d{1,6}", texto(v))) >= 0.9:
        return "num_socio"
    return ""


def _es_texto(v) -> bool:
    t = texto(v)
    return bool(t) and "@" not in t and bool(re.search(r"[^\W\d_]", t))


def sugerir_mapeo(filas, cabecera: bool):
    """Lista con el campo destino de cada columna ('' = no importar)."""
    if not filas:
        return []
    ncols = len(filas[0])
    encabezados = [texto(c) for c in filas[0]] if cabecera else [""] * ncols
    datos = filas[1:] if cabecera else filas
    mapeo = []
    for j in range(ncols):
        col = [f[j] for f in datos]
        contenido = adivinar_por_contenido(col)
        guess = adivinar_por_encabezado(encabezados[j]) if encabezados[j] else ""
        if guess == "num_socio" and contenido != "num_socio":
            guess = ""       # una columna «Socio» con texto no es un número
        mapeo.append(guess or (contenido if not encabezados[j] else ""))

    # Columnas de texto sin cabecera reconocible: nombre, apellidos, dirección
    pendientes = ["nombre", "apellidos", "direccion"]
    usados = set(mapeo)
    for j in range(ncols):
        if mapeo[j] == "" and not encabezados[j]:
            col = [f[j] for f in datos if not es_vacio(f[j])]
            if col and sum(1 for v in col if _es_texto(v)) / len(col) >= 0.6:
                juntos = sum(1 for v in col if len(texto(v).split()) >= 3 or "," in texto(v)) / len(col) >= 0.6
                if juntos and not ({"nombre", "apellidos", "nombre_completo"} & usados):
                    mapeo[j] = "nombre_completo"
                    usados.add("nombre_completo")
                    continue
                for campo in pendientes:
                    if campo not in usados and not (campo in ("nombre", "apellidos") and "nombre_completo" in usados):
                        mapeo[j] = campo
                        usados.add(campo)
                        break
    return mapeo


def etiquetas_columnas(filas, cabecera: bool):
    ncols = len(filas[0]) if filas else 0
    if cabecera and filas:
        return [texto(c) or f"columna {letra(j)}" for j, c in enumerate(filas[0])]
    return [f"columna {letra(j)}" for j in range(ncols)]


# --------------------------------------------------------------------------
# Convertir una fila en datos limpios
# --------------------------------------------------------------------------

def _unir(valores, separador):
    vistos, salida = set(), []
    for v in valores:
        if v and v.casefold() not in vistos:
            vistos.add(v.casefold())
            salida.append(v)
    return separador.join(salida)


def procesar_fila(fila, mapeo, etiquetas, hoy=None):
    hoy = hoy or date.today()
    por_campo = {}
    for j, destino in enumerate(mapeo):
        if destino and j < len(fila):
            por_campo.setdefault(destino, []).append((etiquetas[j], fila[j]))

    def celdas(campo):
        return [(et, v) for et, v in por_campo.get(campo, []) if not es_vacio(v)]

    if not any(celdas(c) for c in por_campo):
        return {"vacia": True}

    avisos, infos, extras = [], [], []
    datos = {c: None for c in CAMPOS_DATOS}

    num = None
    for _et, v in celdas("num_socio"):
        m = re.search(r"\d+", texto(v))
        if m and int(m.group()) >= 1:
            num = num or int(m.group())
        else:
            avisos.append(f"Nº de socio no válido: «{texto(v)}».")

    nombre = _unir([capitalizar(texto(v)) for _e, v in celdas("nombre")], " ")
    apellidos = _unir([capitalizar(texto(v)) for _e, v in celdas("apellidos")], " ")
    dividido = False
    completos = celdas("nombre_completo")
    if completos and (not nombre or not apellidos):
        n2, a2 = dividir_nombre_completo(texto(completos[0][1]))
        dividido = True
        nombre, apellidos = nombre or n2, apellidos or a2
    datos["nombre"], datos["apellidos"] = nombre or None, apellidos or None

    dni = None
    for _e, v in celdas("dni"):
        d, aviso = limpiar_dni(v)
        if d and dni is None:
            dni = d
            if aviso:
                infos.append(aviso)
        elif d and d != dni:
            extras.append(f"otro DNI/NIE en el archivo: {d}")
        elif not d and aviso:
            avisos.append(aviso)
            extras.append(f"DNI/NIE no válido en el archivo: «{texto(v)}»")
    datos["dni"] = dni

    for campo, nacimiento, nombre_campo in (
        ("fecha_nacimiento", True, "Fecha de nacimiento"),
        ("fecha_alta", False, "Fecha de alta"),
    ):
        for _e, v in celdas(campo):
            f = parsear_fecha(v, nacimiento, hoy)
            if f and datos[campo] is None:
                datos[campo] = f
            elif not f:
                avisos.append(f"{nombre_campo} no reconocida: «{texto(v)}».")
                extras.append(f"{nombre_campo.lower()} no reconocida en el archivo: «{texto(v)}»")

    telefonos = []
    for _e, v in celdas("telefono"):
        validos, invalidos = separar_telefonos(v)
        telefonos += [t for t in validos if t not in telefonos]
        for x in invalidos:
            avisos.append(f"Teléfono no válido: «{x}».")
            extras.append(f"teléfono no válido en el archivo: «{x}»")
    if telefonos:
        datos["telefono"] = telefonos[0]
        if len(telefonos) > 1:
            extras.append("otros teléfonos: " + ", ".join(telefonos[1:]))

    correos = []
    for _e, v in celdas("email"):
        validos, invalidos = separar_emails(v)
        correos += [c for c in validos if c not in correos]
        for x in invalidos:
            avisos.append(f"Correo no válido: «{x}».")
            extras.append(f"correo no válido en el archivo: «{x}»")
    if correos:
        datos["email"] = correos[0]
        if len(correos) > 1:
            extras.append("otros correos: " + ", ".join(correos[1:]))

    datos["direccion"] = _unir([texto(v) for _e, v in celdas("direccion")], ", ") or None

    notas = [_unir([texto(v) for _e, v in celdas("notas")], "; ")]
    if extras:
        notas.append("Importación: " + "; ".join(extras) + ".")
    datos["notas"] = "\n".join(n for n in notas if n) or None

    return {
        "vacia": False, "num": num, "datos": datos, "avisos": avisos,
        "infos": infos, "dividido": dividido,
    }


# --------------------------------------------------------------------------
# Decidir qué hacer con un socio que ya existe
# --------------------------------------------------------------------------

def _vacio_campo(campo, valor) -> bool:
    if valor is None:
        return True
    if isinstance(valor, str):
        v = valor.strip()
        return v == "" or (campo == "dni" and v.upper() == DNI_GENERICO)
    return False


def _igual(a, b) -> bool:
    if isinstance(a, str) or isinstance(b, str):
        return str(a or "").strip().casefold() == str(b or "").strip().casefold()
    return a == b


def _txt(v) -> str:
    if v is None:
        return ""
    return v.strftime("%d/%m/%Y") if isinstance(v, date) else str(v).strip()


def instantanea(actual, hoy) -> str:
    partes = [
        f"{ETIQUETA_CAMPO[c]}: {_txt(actual.get(c))}"
        for c in CAMPOS_DATOS if _txt(actual.get(c))
    ]
    return f"[Datos anteriores a la importación del {hoy:%d/%m/%Y}] " + " | ".join(partes)


def fusionar(actual, nuevo, sobrescribir, hoy=None):
    """Devuelve (resultado, campos_cambiados, campos_ignorados)."""
    hoy = hoy or date.today()
    aportados = {
        c: nuevo.get(c) for c in CAMPOS_DATOS if not _vacio_campo(c, nuevo.get(c))
    }
    resultado, cambios, ignorados = dict(actual), [], []

    if not sobrescribir:
        # Prevalece lo que ya hay; solo se rellenan los huecos.
        for c, v in aportados.items():
            if _vacio_campo(c, actual.get(c)):
                resultado[c] = v
                cambios.append(c)
            elif not _igual(actual.get(c), v):
                ignorados.append(c)
        return resultado, cambios, ignorados

    cambios = [c for c, v in aportados.items() if not _igual(actual.get(c), v)]
    if not cambios:
        return resultado, [], []
    foto = instantanea(actual, hoy)
    for c, v in aportados.items():
        if c != "notas":
            resultado[c] = v
    nota = aportados.get("notas") or ""
    resultado["notas"] = (nota + "\n" if nota else "") + foto
    return resultado, cambios, []


# --------------------------------------------------------------------------
# Plan de importación
# --------------------------------------------------------------------------

def _dni_real(dni):
    d = (dni or "").strip().upper()
    return d if d and d != DNI_GENERICO else None


def planificar(filas, mapeo, cabecera, sobrescribir, existentes, hoy=None):
    """Decide qué crear y qué modificar. `existentes`: {num_socio: datos}."""
    hoy = hoy or date.today()
    etiquetas = etiquetas_columnas(filas, cabecera)
    informe = {
        "nuevos": 0, "actualizados": 0, "sin_cambios": 0, "omitidos": 0,
        "leidas": 0, "mensajes": [], "globales": [], "operaciones": [],
    }

    def msg(fila, nivel, t):
        informe["mensajes"].append({"fila": fila, "nivel": nivel, "texto": t})

    estado = {n: dict(d) for n, d in existentes.items()}
    por_dni = {}
    for n, d in estado.items():
        if _dni_real(d.get("dni")):
            por_dni.setdefault(_dni_real(d["dni"]), []).append(n)
    # Números que el propio archivo trae: los asignados automáticamente nunca
    # los pisan y siempre quedan por encima del mayor número existente.
    reservados = set()
    for fila in filas[1 if cabecera else 0:]:
        for j, destino in enumerate(mapeo):
            if destino == "num_socio" and j < len(fila):
                m = re.search(r"\d+", texto(fila[j]))
                if m and int(m.group()) >= 1:
                    reservados.add(int(m.group()))
    tocados = set()
    siguiente = max(list(estado) + list(reservados), default=0) + 1
    divididos = genericos = 0

    for idx in range(1 if cabecera else 0, len(filas)):
        fila_n = idx + 1
        r = procesar_fila(filas[idx], mapeo, etiquetas, hoy)
        if r["vacia"]:
            continue
        informe["leidas"] += 1
        divididos += r["dividido"]
        for a in r["avisos"]:
            msg(fila_n, "aviso", a)
        for i in r["infos"]:
            msg(fila_n, "info", i)
        datos, num = dict(r["datos"]), r["num"]

        existente = None
        if num is not None:
            if num in tocados:
                msg(fila_n, "aviso", f"El nº {num} está repetido en el archivo; se le asigna otro nuevo.")
                num = None
            elif num in estado:
                existente = num
        elif datos["dni"]:
            candidatos = [n for n in por_dni.get(datos["dni"], []) if n not in tocados]
            if len(candidatos) == 1:
                existente = candidatos[0]

        if existente is not None:
            actual = estado[existente]
            if datos["dni"]:
                otros = [n for n in por_dni.get(datos["dni"], []) if n != existente]
                if otros:
                    msg(fila_n, "aviso", f"El DNI {datos['dni']} ya lo tiene el socio #{otros[0]}; no se cambia el DNI de #{existente}.")
                    datos["dni"] = None
            resultado, cambios, ignorados = fusionar(actual, datos, sobrescribir, hoy)
            tocados.add(existente)
            if ignorados:
                msg(fila_n, "info", f"Socio #{existente}: se mantienen los datos actuales de " + ", ".join(ETIQUETA_CAMPO[c].lower() for c in ignorados) + ".")
            if not cambios:
                informe["sin_cambios"] += 1
                continue
            if _dni_real(actual.get("dni")) and existente in por_dni.get(_dni_real(actual["dni"]), []):
                por_dni[_dni_real(actual["dni"])].remove(existente)
            if _dni_real(resultado.get("dni")):
                por_dni.setdefault(_dni_real(resultado["dni"]), []).append(existente)
            estado[existente] = resultado
            informe["actualizados"] += 1
            informe["operaciones"].append({
                "op": "actualizar", "num": existente, "datos": resultado,
                "accion": "Sobrescrito por importación" if sobrescribir else "Completado por importación",
                "detalle": "Campos: " + ", ".join(ETIQUETA_CAMPO[c].lower() for c in cambios) + ".",
            })
            continue

        # ---- socio nuevo ----
        if not datos["nombre"] and not datos["apellidos"]:
            informe["omitidos"] += 1
            msg(fila_n, "error", "Sin nombre ni apellidos: fila omitida.")
            continue
        if not datos["nombre"] or not datos["apellidos"]:
            msg(fila_n, "aviso", "Falta el " + ("nombre" if not datos["nombre"] else "apellido") + "; completar en la ficha.")
        if num is None:
            while siguiente in estado or siguiente in reservados:
                siguiente += 1
            num = siguiente
        dni = datos["dni"]
        if dni and por_dni.get(dni):
            otro = por_dni[dni][0]
            msg(fila_n, "aviso", f"El DNI {dni} ya lo tiene el socio #{otro}; se guarda con el DNI provisional y el original queda en observaciones.")
            nota = f"Importación: DNI {dni} duplicado con el socio #{otro}."
            datos["notas"] = (datos["notas"] + "\n" if datos["notas"] else "") + nota
            dni = None
        if not dni:
            dni = DNI_GENERICO
            genericos += 1
        nuevo = {
            "nombre": datos["nombre"] or "", "apellidos": datos["apellidos"] or "",
            "dni": dni, "fecha_nacimiento": datos["fecha_nacimiento"],
            "telefono": datos["telefono"], "email": datos["email"],
            "direccion": datos["direccion"],
            "fecha_alta": datos["fecha_alta"] or hoy, "notas": datos["notas"],
        }
        estado[num] = nuevo
        tocados.add(num)
        if _dni_real(dni):
            por_dni.setdefault(dni, []).append(num)
        informe["nuevos"] += 1
        informe["operaciones"].append({
            "op": "crear", "num": num, "datos": nuevo,
            "accion": "Alta por importación",
            "detalle": f"Socio importado con el número #{num}.",
        })

    if divididos:
        informe["globales"].append(f"{divididos} nombres completos se dividieron automáticamente en nombre y apellidos. Conviene revisarlos.")
    if genericos:
        informe["globales"].append(f"{genericos} socios nuevos no tenían DNI y se han guardado con el provisional {DNI_GENERICO}.")
    return informe
